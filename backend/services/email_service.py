import base64
import json
import os
import smtplib
import ssl
from datetime import datetime, timezone
from email.message import EmailMessage
from io import BytesIO
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from models.appointment_model import Appointment
from models.contact_model import Contact

NOTIFICATION_RECIPIENT = "cpmittal15@gmail.com"


class EmailNotificationError(Exception):
    """Raised when a record notification cannot be configured or delivered."""


def _create_workbook():
    workbook = Workbook()
    contacts_sheet = workbook.active
    contacts_sheet.title = "Inquiries"
    appointments_sheet = workbook.create_sheet("Appointments")

    sheets = [
        (
            contacts_sheet,
            ("ID", "Name", "Email", "Phone", "Subject", "Message"),
            Contact.query.order_by(Contact.id.asc()).all(),
            ("id", "name", "email", "phone", "subject", "message"),
        ),
        (
            appointments_sheet,
            ("ID", "Name", "Email", "Phone", "Service", "Date", "Message"),
            Appointment.query.order_by(Appointment.id.asc()).all(),
            ("id", "name", "email", "phone", "service", "date", "message"),
        ),
    ]

    for sheet, headings, records, fields in sheets:
        sheet.append(headings)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1D4ED8")

        for record in records:
            values = [getattr(record, field) or "" for field in fields]
            sheet.append(values)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = max(
                len(str(cell.value or "")) for cell in column
            )
            sheet.column_dimensions[
                get_column_letter(column[0].column)
            ].width = min(max(width + 2, 12), 60)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def send_email(subject, body):
    workbook = _create_workbook()
    filename = (
        f"portfolio-records-{datetime.now(timezone.utc):%Y%m%d-%H%M%S}.xlsx"
    )
    resend_api_key = os.environ.get("RESEND_API_KEY")
    if resend_api_key:
        sender = os.environ.get("EMAIL_FROM")
        if not sender:
            raise EmailNotificationError(
                "EMAIL_FROM must be configured when using Resend."
            )

        payload = {
            "from": sender,
            "to": [NOTIFICATION_RECIPIENT],
            "subject": subject,
            "text": body,
            "attachments": [
                {
                    "filename": filename,
                    "content": base64.b64encode(workbook).decode("ascii"),
                    "content_type": (
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                }
            ],
        }
        request = Request(
            "https://api.resend.com/emails",
            data=json.dumps(payload).encode("utf-8"),
            headers={
                "Authorization": f"Bearer {resend_api_key}",
                "Content-Type": "application/json",
            },
            method="POST",
        )
        try:
            with urlopen(request, timeout=20):
                return
        except (HTTPError, URLError, TimeoutError, OSError) as error:
            raise EmailNotificationError("Resend API delivery failed.") from error

    username = os.environ.get("SMTP_USERNAME")
    password = os.environ.get("SMTP_PASSWORD")
    if not username or not password:
        raise EmailNotificationError(
            "Configure RESEND_API_KEY and EMAIL_FROM, or SMTP_USERNAME and "
            "SMTP_PASSWORD, to enable email notifications."
        )

    host = os.environ.get("SMTP_HOST", "smtp.gmail.com")
    try:
        port = int(os.environ.get("SMTP_PORT", "587"))
    except ValueError as error:
        raise EmailNotificationError(
            "SMTP_PORT must be a valid port number."
        ) from error

    message = EmailMessage()
    message["Subject"] = subject
    message["From"] = username
    message["To"] = NOTIFICATION_RECIPIENT
    message.set_content(body)
    message.add_attachment(
        workbook,
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=filename,
    )

    try:
        with smtplib.SMTP(host, port, timeout=20) as smtp:
            smtp.ehlo()
            smtp.starttls(context=ssl.create_default_context())
            smtp.ehlo()
            smtp.login(username, password)
            smtp.send_message(message)
    except (OSError, smtplib.SMTPException) as error:
        raise EmailNotificationError("SMTP delivery failed.") from error
